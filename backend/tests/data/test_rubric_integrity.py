"""
Rubric data integrity: every rubric document the graders depend on exists, parses, and the
subject dropdown (/api/ocr/subjects) only offers subjects the grading pipeline can resolve.
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

import pytest
from docx import Document

from backend.utils import rubric_loader
from support.env import REPO_ROOT

BACKEND = REPO_ROOT / "backend"
DROPDOWN_ROOT = BACKEND / "20marks_Rubrics"  # read by utils.rubric_loader (subject list)
GRADER_ROOT = BACKEND / "ocr" / "20marks_Rubrics"  # read by ocr.grade_pdf_answer (actual grading)

SUBJECTS = rubric_loader.list_subject_rubrics()
SUBJECT_IDS = [s.subject_id for s in SUBJECTS]

# Rubric/asset files referenced by name from the essay, outline and precis pipelines.
PIPELINE_ASSETS = [
    "eng_essay/CSS English Essay Evaluation Rubric Based on FPSC Examiners.docx",
    "eng_essay/ANNOTATIONS RUBRIC FOR ESSAY.docx",
    "eng_essay/Report Format.docx",
    "outline/CSS English Essay Outline Evaluation Rubric Based on FPSC Examiners.docx",
    "outline/ANNOTATIONS FOR ESSAY OUTLINE.docx",
    "precis/Precis Rubric.docx",
    "precis/ANNOTATIONS RUBRIC FOR PRECIS.docx",
    "precis/colouring_scheme.jpeg",
    "ocr/REFINED RUBRIC.docx",
    "ocr/fonts/NotoSans-Regular.ttf",
    "eng_essay/fonts/NotoSans-Regular.ttf",
]


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(t for t in parts if t.strip())


def _tree_digest(root: Path) -> dict[str, str]:
    return {
        p.relative_to(root).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in sorted(root.rglob("*.docx"))
    }


# ---------------------------------------------------------------- subject list


@pytest.mark.unit
def test_subject_list_is_populated_and_well_formed():
    assert len(SUBJECTS) >= 20
    assert len(set(SUBJECT_IDS)) == len(SUBJECT_IDS), "two folders normalise to the same subject id"
    for subject in SUBJECTS:
        assert re.fullmatch(r"[a-z0-9]+(-[a-z0-9]+)*", subject.subject_id), subject.subject_id
        assert subject.doc_path.is_file()
        assert subject.doc_path.parent.name == subject.display_name


@pytest.mark.unit
@pytest.mark.parametrize("subject_id", SUBJECT_IDS)
def test_every_subject_rubric_parses_to_real_text(subject_id):
    rubric_loader.load_rubric_text.cache_clear()
    text = rubric_loader.load_rubric_text(subject_id)
    assert len(text) > 500, f"{subject_id} rubric looks empty ({len(text)} chars)"


@pytest.mark.unit
@pytest.mark.parametrize(
    ("given", "expected"),
    [("Political Science", "political-science"), ("  History_of  Indo-Pak ", "history-of-indo-pak"), ("US History!", "us-history")],
)
def test_subject_name_normalisation(given, expected):
    assert rubric_loader._normalize_subject_name(given) == expected


@pytest.mark.unit
def test_unknown_subject_error_lists_available_ids():
    with pytest.raises(FileNotFoundError, match="political-science"):
        rubric_loader.load_rubric_text("astrology")


@pytest.mark.unit
def test_display_name_lookup():
    assert rubric_loader.get_subject_display_name("political-science") == "Political Science"
    assert rubric_loader.get_subject_display_name("not-a-subject") == "not-a-subject"


# ---------------------------------------------------------------- dropdown vs grader


@pytest.mark.api
def test_subjects_endpoint_matches_rubric_folders(client):
    body = client.get("/api/ocr/subjects").json()
    assert body["count"] == len(SUBJECTS)
    assert body["subjects"] == [{"id": s.subject_id, "display_name": s.display_name} for s in SUBJECTS]
    assert isinstance(body["latency_ms"], int)


@pytest.mark.api
def test_regular_subjects_endpoint_is_non_empty(client):
    body = client.get("/api/ocr-regular/subjects").json()
    assert body["count"] == len(body["subjects"]) > 0


@pytest.mark.unit
def test_grader_rubric_tree_mirrors_dropdown_tree():
    """The dropdown lists backend/20marks_Rubrics but grading reads backend/ocr/20marks_Rubrics."""
    assert _tree_digest(DROPDOWN_ROOT) == _tree_digest(GRADER_ROOT)


@pytest.mark.unit
@pytest.mark.parametrize("subject_id", SUBJECT_IDS)
def test_grader_resolves_every_dropdown_subject(subject_id):
    from backend.ocr.grade_pdf_answer import find_subject_rubric_path

    path = find_subject_rubric_path(subject_id)
    assert path is not None, f"grader cannot find a rubric for '{subject_id}'"
    display = next(s.display_name for s in SUBJECTS if s.subject_id == subject_id)
    assert Path(path).parent.name == display


@pytest.mark.unit
def test_grader_returns_none_for_unknown_subject():
    from backend.ocr.grade_pdf_answer import find_subject_rubric_path

    assert find_subject_rubric_path("astrology") is None
    assert find_subject_rubric_path("") is None


# ---------------------------------------------------------------- other pipelines


@pytest.mark.unit
@pytest.mark.parametrize("relpath", PIPELINE_ASSETS)
def test_pipeline_assets_exist(relpath):
    path = BACKEND / relpath
    assert path.is_file(), f"missing {relpath}"
    assert path.stat().st_size > 0
    if path.suffix == ".docx":
        assert len(_docx_text(path)) > 100
